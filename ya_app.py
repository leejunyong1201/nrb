import os
import re
import socket
import openpyxl
import streamlit as st
import streamlit.components.v1 as components

st.set_page_config(page_title="야적장 제품 위치 검색 시스템", layout="wide")

# 로컬 IP 주소 가져오기
def get_local_ip():
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(("8.8.8.8", 80))
        ip = s.getsockname()[0]
        s.close()
        return ip
    except Exception:
        return "127.0.0.1"

local_ip = get_local_ip()

# 헤더 및 안내
st.title("📦 야적장 제품 위치 검색 시스템")

# 상단 접속 링크 안내 구역
st.info(f"""
🌐 **현재 접속 및 공유 정보**
- **이 PC에서 접속:** `http://localhost:8501` (또는 지정 포트)
- **같은 와이파이/사내망 접속:** `http://{local_ip}:8501`
""")

st.markdown("---")

# 1. 엑셀 파일 업로드/선택 기능 (data.xlsx 자동 로드 반영)
uploaded_file = st.sidebar.file_uploader("📂 엑셀 파일(.xlsx) 선택 (선택 사항)", type=["xlsx"])

EXCEL_FILE = None

if uploaded_file is not None:
    temp_dir = os.path.join(os.path.expanduser("~"), "Downloads")
    os.makedirs(temp_dir, exist_ok=True)
    EXCEL_FILE = os.path.join(temp_dir, "temp_yard_layout.xlsx")
    with open(EXCEL_FILE, "wb") as f:
        f.write(uploaded_file.getbuffer())
    st.sidebar.success(f"✅ 업로드 파일 적용: {uploaded_file.name}")

elif os.path.exists("data.xlsx"):
    EXCEL_FILE = "data.xlsx"
    st.sidebar.info("ℹ️ GitHub 기본 파일(data.xlsx) 로드 완료")

else:
    default_path = os.path.join(os.path.expanduser("~"), "Downloads", "12321341234.xlsx")
    if os.path.exists(default_path):
        EXCEL_FILE = default_path
        st.sidebar.info("ℹ️ 로컬 기본 엑셀 파일을 불러왔습니다.")

def get_hex_color(color_obj):
    if not color_obj:
        return None
    if hasattr(color_obj, 'rgb') and color_obj.rgb:
        rgb = color_obj.rgb
        if isinstance(rgb, str):
            if len(rgb) == 8:
                return f"#{rgb[2:]}"
            elif len(rgb) == 6:
                return f"#{rgb}"
    return None

@st.cache_data(ttl=300)
def load_excel_full_style(file_path):
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.active
    
    matrix = []
    merged_cells = ws.merged_cells.ranges
    skip_coords = set()

    max_r = ws.max_row
    max_c = ws.max_column

    for r in range(1, max_r + 1):
        row_data = []
        for c in range(1, max_c + 1):
            if (r, c) in skip_coords:
                continue

            cell = ws.cell(row=r, column=c)
            val = str(cell.value) if cell.value is not None else ""
            
            bg_color = None
            if cell.fill and cell.fill.fill_type:
                extracted = get_hex_color(cell.fill.start_color)
                if extracted and extracted.upper() not in ["#000000", "#FFFFFF"]:
                    bg_color = extracted
            
            fg_color = None
            is_bold = False
            if cell.font:
                is_bold = bool(cell.font.bold)
                fg_color = get_hex_color(cell.font.color)

            rowspan, colspan = 1, 1
            for m in merged_cells:
                if r == m.min_row and c == m.min_col:
                    rowspan = m.max_row - m.min_row + 1
                    colspan = m.max_col - m.min_col + 1
                    for mr in range(m.min_row, m.max_row + 1):
                        for mc in range(m.min_col, m.max_col + 1):
                            skip_coords.add((mr, mc))
                    break

            row_data.append({
                'r': r,
                'c': c,
                'val': val,
                'bg_color': bg_color,
                'fg_color': fg_color,
                'is_bold': is_bold,
                'rowspan': rowspan,
                'colspan': colspan
            })
        if row_data:
            matrix.append(row_data)
            
    return matrix

def generate_possible_hyphen_patterns(digits):
    patterns = []
    if len(digits) >= 5 and digits[:3].isdigit():
        dong = digits[:3]
        rest = digits[3:]
        
        if len(rest) == 2:
            patterns.append(f"{dong}-{rest[0]}-{rest[1]}")
        elif len(rest) == 3:
            f_a, n_a = rest[:2], rest[2:]
            if 1 <= int(n_a) <= 14:
                patterns.append(f"{dong}-{f_a}-{n_a}")
                
            f_b, n_b = rest[:1], rest[1:]
            if 1 <= int(n_b) <= 14:
                patterns.append(f"{dong}-{f_b}-{n_b}")
        elif len(rest) == 4:
            f_c, n_c = rest[:2], rest[2:]
            if 1 <= int(n_c) <= 14:
                patterns.append(f"{dong}-{f_c}-{n_c}")
                
    return patterns

def is_exact_match(search_query, cell_text):
    if not search_query or not cell_text:
        return False
    
    clean_query = search_query.strip()
    if not clean_query:
        return False

    cell_str = str(cell_text).strip()

    if clean_query.lower() == cell_str.lower():
        return True

    target_query = clean_query
    if " " in clean_query:
        target_query = "-".join(re.split(r'\s+', clean_query))

    if "-" in target_query:
        escaped_query = re.escape(target_query)
        pattern = r'(?<![0-9a-zA-Z가-힣_])' + escaped_query + r'(?![0-9a-zA-Z가-힣_])'
        if re.search(pattern, cell_str, re.IGNORECASE):
            return True

    if clean_query.isdigit():
        possible_hyphen_formats = generate_possible_hyphen_patterns(clean_query)
        for fmt in possible_hyphen_formats:
            escaped_fmt = re.escape(fmt)
            pattern = r'(?<![0-9a-zA-Z가-힣_])' + escaped_fmt + r'(?![0-9a-zA-Z가-힣_])'
            if re.search(pattern, cell_str, re.IGNORECASE):
                return True

    tokens = re.split(r'[\s/,\n\r]+', cell_str)
    for token in tokens:
        if clean_query.lower() == token.strip().lower():
            return True

    return False

# 메인 UI 및 화면 출력 부분
if not EXCEL_FILE or not os.path.exists(EXCEL_FILE):
    st.warning("👈 등록된 엑셀 파일이 없습니다. GitHub 저장소에 'data.xlsx' 파일을 업로드하거나, 왼쪽 사이드바에서 배치도 엑셀(.xlsx) 파일을 선택해 주세요.")
else:
    col1, col2 = st.columns([3, 1])
    with col1:
        search_term = st.text_input("🔍 검색할 제품명, 제품번호 또는 구역을 입력하세요 (예: 401 11 1 또는 401-11-1)", "").strip()
    with col2:
        zoom_level = st.slider("🔍 배치도 크기 조절 (%)", min_value=20, max_value=100, value=50, step=5)

    matrix = load_excel_full_style(EXCEL_FILE)

    html = [f'''
    <!DOCTYPE html>
    <html>
    <head>
    <meta name="viewport" content="width=device-width, initial-scale=1.0, maximum-scale=50.0, user-scalable=yes">
    <style>
        body {{
            font-family: 'Malgun Gothic', 'Segoe UI', sans-serif;
            margin: 0;
            padding: 5px;
            background-color: #ffffff;
            touch-action: none;
            overflow: auto;
        }}
        .viewport-wrapper {{
            overflow: auto;
            max-width: 100%;
            max-height: 800px;
            position: relative;
        }}
        .table-container {{
            transform-origin: 0 0;
            zoom: {zoom_level}%;
            display: inline-block;
            transition: transform 0.05s ease-out;
        }}
        table {{
            border-collapse: collapse;
            font-size: 11px;
            text-align: center;
        }}
        td {{
            border: 1px solid #333333;
            padding: 2px 4px;
            white-space: pre-line;
            word-break: break-all;
            min-width: 35px;
            height: 22px;
            line-height: 1.25;
            vertical-align: middle;
        }}
        .matched {{
            background-color: #FFFF00 !important;
            color: #000000 !important;
            font-weight: bold !important;
            font-size: 12px !important;
            border: 3px solid #FF0000 !important;
            box-shadow: 0 0 8px #FF0000;
        }}
    </style>
    </head>
    <body>
    <div class="viewport-wrapper" id="viewport">
        <div class="table-container" id="target">
            <table>
    ''']

    matched_items = []

    for row in matrix:
        html.append('<tr>')
        for cell in row:
            val = cell['val']
            r, c = cell['r'], cell['c']
            
            is_matched = is_exact_match(search_term, val)
            if is_matched:
                matched_items.append((r, c, val))

            styles = []
            if cell['bg_color'] and not is_matched:
                styles.append(f"background-color: {cell['bg_color']};")
            if cell['fg_color'] and not is_matched:
                styles.append(f"color: {cell['fg_color']};")
            if cell['is_bold']:
                styles.append("font-weight: bold;")

            style_attr = f' style="{" ".join(styles)}"' if styles else ''
            class_attr = ' class="matched"' if is_matched else ''
            span_attr = f' rowspan="{cell["rowspan"]}"' if cell['rowspan'] > 1 else ''
            span_attr += f' colspan="{cell["colspan"]}"' if cell['colspan'] > 1 else ''

            html.append(f'<td{span_attr}{class_attr}{style_attr}>{val}</td>')
        html.append('</tr>')

    html.append('''
            </table>
        </div>
    </div>

    <!-- 모바일 터치 손가락 줌(Pinch-to-zoom) 스크립트 -->
    <script>
        const viewport = document.getElementById('viewport');
        const target = document.getElementById('target');
        
        let currentScale = 1;
        let startDist = 0;
        let initialScale = 1;

        viewport.addEventListener('touchstart', (e) => {
            if (e.touches.length === 2) {
                e.preventDefault();
                startDist = Math.hypot(
                    e.touches[0].pageX - e.touches[1].pageX,
                    e.touches[0].pageY - e.touches[1].pageY
                );
                initialScale = currentScale;
            }
        }, { passive: false });

        viewport.addEventListener('touchmove', (e) => {
            if (e.touches.length === 2) {
                e.preventDefault();
                const dist = Math.hypot(
                    e.touches[0].pageX - e.touches[1].pageX,
                    e.touches[0].pageY - e.touches[1].pageY
                );
                if (startDist > 0) {
                    const factor = dist / startDist;
                    currentScale = Math.min(Math.max(initialScale * factor, 0.5), 50.0);
                    target.style.transform = `scale(${currentScale})`;
                }
            }
        }, { passive: false });
    </script>
    </body>
    </html>
    ''')

    st.subheader(f"🗺 야적장 원본 전체 배치도 (기본 배율: {zoom_level}%)")
    components.html("".join(html), height=850, scrolling=True)

    if search_term:
        st.markdown("---")
        if matched_items:
            st.success(f"🎯 총 **{len(matched_items)}건**의 위치를 찾았습니다!")
            st.subheader("📋 검색 상세 위치 결과")
            res_data = [{"행 위치": f"{item[0]}행", "열 위치": f"{item[1]}열", "내용": item[2]} for item in matched_items]
            st.table(res_data)
        else:
            st.warning(f"⚠️ '{search_term}'에 대한 검색 결과가 없습니다.")
